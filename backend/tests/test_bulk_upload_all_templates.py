import os
import io
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.db.session import SessionLocal
from app.models.employee import CompanyEmployee
from app.models.expense import CompanyExpense
from app.models.product import CompanyProduct
from app.models.indus_gbpa import IndusCustomerGbpa
from app.models.indus_site import IndusSiteDetails
from app.models.indus_infra import IndusCustomerInfra
from app.models.indus_gbpa_material import IndusCustomerGbpaMaterial
from app.models.indus_gbpa_expense import IndusCustomerGbpaExpense
from sqlalchemy import or_

client = TestClient(app)

TPL_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "frontend", "bulk upload template")

def create_populated_wb(filename):
    src_path = os.path.join(TPL_DIR, filename)
    wb = openpyxl.load_workbook(src_path)
    ws = wb.active
    
    fn = filename.lower()
    if "employee" in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value="Staff")
            ws.cell(row=i, column=2, value=f"Test Employee {i}")
            ws.cell(row=i, column=3, value=f"TEST_BULK_EMP_{i}")
            ws.cell(row=i, column=4, value="123 Test Street")
            ws.cell(row=i, column=5, value=f"987654321{i}")
            ws.cell(row=i, column=6, value=f"test{i}@example.com")
            ws.cell(row=i, column=7, value="15/05/1990")
            ws.cell(row=i, column=8, value="O+")
            ws.cell(row=i, column=9, value="Single")
            ws.cell(row=i, column=10, value="01/01/2022")
            ws.cell(row=i, column=11, value="Engineer")
            ws.cell(row=i, column=12, value="2 Years")
            ws.cell(row=i, column=13, value="B.Tech")
            ws.cell(row=i, column=14, value=f"ABCDE123{i}F")
            ws.cell(row=i, column=15, value=f"12345678901{i}")
            ws.cell(row=i, column=20, value="Yes")
            ws.cell(row=i, column=21, value="Active")
    elif "expense" in fn and "gbpa" not in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value="Capex")
            ws.cell(row=i, column=2, value="Travel")
            ws.cell(row=i, column=3, value="Hotel")
            ws.cell(row=i, column=4, value=f"TEST_EXP_{i}")
            ws.cell(row=i, column=5, value=f"Bulk Test Expense {i}")
            ws.cell(row=i, column=6, value="Nos")
            ws.cell(row=i, column=7, value="18%")
            ws.cell(row=i, column=8, value="10%")
            ws.cell(row=i, column=9, value="No")
    elif "product" in fn and "gbpa" not in fn and "infra" not in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value=f"Test Mat Head {i}")
            ws.cell(row=i, column=2, value="Tower Infrastructure")
            ws.cell(row=i, column=3, value=f"TEST_PRD_{i}")
            ws.cell(row=i, column=4, value=f"Bulk Test Product {i}")
            ws.cell(row=i, column=5, value="Nos")
            ws.cell(row=i, column=6, value="Nos")
            ws.cell(row=i, column=7, value="18%")
            ws.cell(row=i, column=8, value="1.0")
            ws.cell(row=i, column=9, value="100")
            ws.cell(row=i, column=10, value="50")
            ws.cell(row=i, column=11, value="5%")
            ws.cell(row=i, column=12, value="2%")
            ws.cell(row=i, column=13, value="Active")
    elif "gbpa" in fn and "material" not in fn and "expense" not in fn and "infra" not in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value="Indus Tower Ltd")
            ws.cell(row=i, column=2, value=f"TEST_GBPA_{i}")
            ws.cell(row=i, column=3, value=f"GBPA Item {i}")
            ws.cell(row=i, column=4, value="Capex")
            ws.cell(row=i, column=5, value=f"GBPA Desc {i}")
            ws.cell(row=i, column=6, value="Pcs")
            ws.cell(row=i, column=7, value=1500.0)
            ws.cell(row=i, column=8, value="HSN")
            ws.cell(row=i, column=9, value="73082019")
            ws.cell(row=i, column=10, value="10%")
            ws.cell(row=i, column=11, value=150.0)
    elif "site_details" in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value="Indus Tower Ltd")
            ws.cell(row=i, column=2, value=f"TEST_SITE_{i}")
            ws.cell(row=i, column=3, value=f"WH_{i}")
            ws.cell(row=i, column=4, value=f"Bulk Site Name {i}")
            ws.cell(row=i, column=5, value="GBT")
            ws.cell(row=i, column=6, value="Central")
            ws.cell(row=i, column=7, value="City Center")
            ws.cell(row=i, column=8, value="Site Address 123")
            ws.cell(row=i, column=9, value=12.9716)
            ws.cell(row=i, column=10, value=77.5946)
            ws.cell(row=i, column=11, value="North Zone")
            ws.cell(row=i, column=12, value="Engineer Ramesh")
            ws.cell(row=i, column=13, value="Manager Suresh")
    elif "site_contacts" in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value="Indus Tower Ltd")
            ws.cell(row=i, column=2, value=f"TEST_SITE_{i}")
            ws.cell(row=i, column=3, value=f"FSE Contact {i}")
            ws.cell(row=i, column=4, value=f"AOM Contact {i}")
    elif "infra products" in fn or "indus infra" in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value="Indus Tower Ltd")
            ws.cell(row=i, column=2, value=f"TEST_INFRA_{i}")
            ws.cell(row=i, column=3, value="Diesel Generator")
            ws.cell(row=i, column=4, value=f"Generator 25kVA {i}")
            ws.cell(row=i, column=5, value="Kirloskar")
            ws.cell(row=i, column=6, value="Nos")
            ws.cell(row=i, column=7, value="Yes")
            ws.cell(row=i, column=8, value="Yes")
    elif "materials" in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value="Indus Tower Ltd")
            ws.cell(row=i, column=2, value=f"GBPA_ITM_{i}")
            ws.cell(row=i, column=3, value=f"Material Head {i}")
            ws.cell(row=i, column=4, value="Cables")
            ws.cell(row=i, column=5, value="Parent")
            ws.cell(row=i, column=6, value=f"TEST_MAT_{i}")
            ws.cell(row=i, column=7, value=f"Material Desc {i}")
    elif "expenses" in fn:
        for i in range(2, 5):
            ws.cell(row=i, column=1, value="Indus Tower Ltd")
            ws.cell(row=i, column=2, value=f"GBPA_ITM_{i}")
            ws.cell(row=i, column=3, value=f"Expense Head {i}")
            ws.cell(row=i, column=4, value="Logistics")
            ws.cell(row=i, column=5, value="Parent")
            ws.cell(row=i, column=6, value=f"TEST_GEXP_{i}")
            ws.cell(row=i, column=7, value=f"Expense Desc {i}")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

@pytest.fixture(autouse=True)
def cleanup_test_records():
    yield
    db = SessionLocal()
    try:
        db.query(CompanyEmployee).filter(CompanyEmployee.employee_code.like("TEST_BULK_%")).delete(synchronize_session=False)
        db.query(CompanyExpense).filter(CompanyExpense.expense_code.like("TEST_EXP_%")).delete(synchronize_session=False)
        db.query(CompanyProduct).filter(or_(CompanyProduct.material_code.like("TEST_PRD_%"), CompanyProduct.material_head.like("Test Mat Head%"))).delete(synchronize_session=False)
        db.query(IndusCustomerGbpa).filter(IndusCustomerGbpa.item_code.like("TEST_GBPA_%")).delete(synchronize_session=False)
        db.query(IndusSiteDetails).filter(IndusSiteDetails.site_code.like("TEST_SITE_%")).delete(synchronize_session=False)
        db.query(IndusCustomerInfra).filter(IndusCustomerInfra.item_code.like("TEST_INFRA_%")).delete(synchronize_session=False)
        db.query(IndusCustomerGbpaMaterial).filter(IndusCustomerGbpaMaterial.material_code.like("TEST_MAT_%")).delete(synchronize_session=False)
        db.query(IndusCustomerGbpaExpense).filter(IndusCustomerGbpaExpense.expense_code.like("TEST_GEXP_%")).delete(synchronize_session=False)
        db.commit()
    finally:
        db.close()

TEST_CONFIGS = [
    ("Company_Employee-Template.xlsx", "/api/v1/master/employees/bulk-upload"),
    ("Company_Expenses-Template.xlsx", "/api/v1/master/expenses/bulk-upload"),
    ("Company_Products-Template.xlsx", "/api/v1/master/products/bulk-upload"),
    ("Customer_GBPA-Template.xlsx", "/api/v1/customer/indus/gbpa/bulk-upload"),
    ("Site_details-Template.xlsx", "/api/v1/customer/indus/sites/bulk-upload"),
    ("Site_contacts-Template.xlsx", "/api/v1/customer/indus/sites/bulk-upload"),
    ("indus Infra Products-Template.xlsx", "/api/v1/customer/indus/infra/bulk-upload"),
    ("indus GBPA Products-Materials-Template.xlsx", "/api/v1/customer/indus/gbpa/materials/bulk-upload"),
    ("indus GBPA Products-Expenses-Template.xlsx", "/api/v1/customer/indus/gbpa/expenses/bulk-upload"),
]

@pytest.mark.parametrize("filename,endpoint", TEST_CONFIGS)
def test_bulk_upload_workflow_and_validation(filename, endpoint):
    content = create_populated_wb(filename)
    
    # 1. Valid upload
    res = client.post(endpoint, files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert res.status_code == 200, f"Failed upload for {filename}: {res.text}"
    data = res.json()
    assert data.get("success") is True
    assert data.get("imported_count") is not None

    # 2. Idempotent Retry
    res_retry = client.post(endpoint, files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert res_retry.status_code == 200
    retry_data = res_retry.json()
    assert retry_data.get("success") is True

    # 3. Mismatched template rejection
    mismatched_filename = "Company_Employee-Template.xlsx" if filename != "Company_Employee-Template.xlsx" else "indus Infra Products-Template.xlsx"
    mismatched_content = create_populated_wb(mismatched_filename)
    res_wrong = client.post(endpoint, files={"file": (mismatched_filename, mismatched_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert res_wrong.status_code == 400
    assert "Invalid" in res_wrong.json().get("detail", "")
