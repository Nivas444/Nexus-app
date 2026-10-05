import math
from decimal import Decimal
from typing import Optional, Dict, Any
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.repositories.indus_gbpa_repository import IndusGbpaRepository
from app.models.indus_gbpa import IndusCustomerGbpa
from app.schemas.indus_gbpa import (
    IndusGbpaCreate,
    IndusGbpaUpdate,
    IndusGbpaResponse,
    IndusGbpaListResponse
)

def _parse_numeric(val: Any) -> Optional[Decimal]:
    if val is None:
        return None
    if isinstance(val, (int, float, Decimal)):
        return Decimal(str(val))
    if isinstance(val, str):
        cleaned = val.replace(",", "").replace("%", "").replace("₹", "").strip()
        if not cleaned:
            return None
        try:
            return Decimal(cleaned)
        except Exception:
            return None
    return None

def _format_rate(val: Any) -> str:
    if val is None:
        return "0.00"
    try:
        return f"{float(val):.2f}"
    except Exception:
        return str(val)

def _format_gbpa_response(item: IndusCustomerGbpa) -> IndusGbpaResponse:
    rate_str = _format_rate(item.rate)
    budget_pct_str = f"{float(item.budget_percentage):.0f}" if item.budget_percentage is not None else ""
    budget_amt_str = _format_rate(item.budget_amount) if item.budget_amount is not None else ""

    return IndusGbpaResponse(
        item_id=item.item_id,
        id=item.item_id,
        company_name=item.company_name,
        customer_name=item.customer_name,
        item_code=item.item_code,
        itemCode=item.item_code,
        item_name=item.item_name,
        itemName=item.item_name,
        productName=item.item_name,
        item_description=item.item_description,
        itemDescription=item.item_description,
        productDescription=item.item_description,
        item_type=item.item_type,
        itemType=item.item_type,
        productType=item.item_type,
        hsn_sac=item.hsn_sac,
        hsnSacType=item.hsn_sac,
        hsn_sac_code=item.hsn_sac_code,
        hsnSacCode=item.hsn_sac_code,
        uom=item.uom or "Pcs",
        rate=rate_str,
        activeRate=rate_str,
        gst_rate=item.gst_rate,
        budget_percentage=item.budget_percentage,
        budgetPercent=budget_pct_str,
        budget_amount=item.budget_amount,
        budgetAmount=budget_amt_str,
        status=item.status or "Active",
        created_at=item.created_at,
        updated_at=item.updated_at
    )

class IndusGbpaService:
    def __init__(self, db: Session):
        self.db = db
        self.repo = IndusGbpaRepository(db)

    def get_gbpa_list(
        self,
        page: int = 1,
        page_size: int = 100,
        search: Optional[str] = None,
        item_type: Optional[str] = None,
        status_filter: Optional[str] = None,
        sort_by: str = "item_id",
        sort_desc: bool = False
    ) -> IndusGbpaListResponse:
        skip = (page - 1) * page_size
        items, total = self.repo.list_items(
            skip=skip,
            limit=page_size,
            search=search,
            item_type=item_type,
            status_filter=status_filter,
            sort_by=sort_by,
            sort_desc=sort_desc
        )

        pages = math.ceil(total / page_size) if total > 0 else 1
        formatted_items = [_format_gbpa_response(item) for item in items]

        return IndusGbpaListResponse(
            items=formatted_items,
            total=total,
            page=page,
            page_size=page_size,
            pages=pages
        )

    def get_gbpa_by_id(self, item_id: int) -> IndusGbpaResponse:
        item = self.repo.get_by_id(item_id)
        if not item:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"GBPA Item with ID {item_id} not found."
            )
        return _format_gbpa_response(item)

    def create_gbpa(self, payload: IndusGbpaCreate) -> IndusGbpaResponse:
        # Check uniqueness of item_code if provided
        if payload.item_code and payload.item_code.strip():
            existing_code = self.repo.get_by_item_code(payload.item_code.strip())
            if existing_code:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"GBPA Item Code '{payload.item_code.strip()}' already exists for Indus Tower Ltd."
                )

        # Check uniqueness of item_name if provided
        if payload.item_name and payload.item_name.strip():
            existing_name = self.repo.get_by_item_name(payload.item_name.strip())
            if existing_name:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"GBPA Item Name '{payload.item_name.strip()}' already exists for Indus Tower Ltd."
                )

        item_dict = {
            "company_name": payload.company_name or "Nexus",
            "customer_name": payload.customer_name or "Indus Tower Ltd",
            "item_code": payload.item_code.strip() if payload.item_code else None,
            "item_name": payload.item_name.strip() if payload.item_name else None,
            "item_description": payload.item_description.strip() if payload.item_description else None,
            "item_type": payload.item_type.strip() if payload.item_type else "Capex",
            "hsn_sac": payload.hsn_sac.strip() if payload.hsn_sac else "HSN",
            "hsn_sac_code": payload.hsn_sac_code.strip() if payload.hsn_sac_code else None,
            "uom": payload.uom.strip() if payload.uom else "Pcs",
            "rate": _parse_numeric(payload.rate),
            "gst_rate": _parse_numeric(payload.gst_rate),
            "budget_percentage": _parse_numeric(payload.budget_percentage),
            "budget_amount": _parse_numeric(payload.budget_amount),
            "status": payload.status.strip() if payload.status else "Active"
        }

        created = self.repo.create(item_dict)
        return _format_gbpa_response(created)

    def update_gbpa(self, item_id: int, payload: IndusGbpaUpdate) -> IndusGbpaResponse:
        item = self.repo.get_by_id(item_id)
        if not item:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"GBPA Item with ID {item_id} not found."
            )

        # Check uniqueness of item_code if changed
        if payload.item_code and payload.item_code.strip():
            existing_code = self.repo.get_by_item_code(payload.item_code.strip())
            if existing_code and existing_code.item_id != item_id:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"GBPA Item Code '{payload.item_code.strip()}' is already in use by another record."
                )

        # Check uniqueness of item_name if changed
        if payload.item_name and payload.item_name.strip():
            existing_name = self.repo.get_by_item_name(payload.item_name.strip())
            if existing_name and existing_name.item_id != item_id:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"GBPA Item Name '{payload.item_name.strip()}' is already in use by another record."
                )

        update_dict: Dict[str, Any] = {}
        if payload.company_name is not None:
            update_dict["company_name"] = payload.company_name.strip()
        if payload.customer_name is not None:
            update_dict["customer_name"] = payload.customer_name.strip()
        if payload.item_code is not None:
            update_dict["item_code"] = payload.item_code.strip() if payload.item_code else None
        if payload.item_name is not None:
            update_dict["item_name"] = payload.item_name.strip() if payload.item_name else None
        if payload.item_description is not None:
            update_dict["item_description"] = payload.item_description.strip() if payload.item_description else None
        if payload.item_type is not None:
            update_dict["item_type"] = payload.item_type.strip() if payload.item_type else None
        if payload.hsn_sac is not None:
            update_dict["hsn_sac"] = payload.hsn_sac.strip() if payload.hsn_sac else None
        if payload.hsn_sac_code is not None:
            update_dict["hsn_sac_code"] = payload.hsn_sac_code.strip() if payload.hsn_sac_code else None
        if payload.uom is not None:
            update_dict["uom"] = payload.uom.strip() if payload.uom else None
        if payload.rate is not None:
            update_dict["rate"] = _parse_numeric(payload.rate)
        if payload.gst_rate is not None:
            update_dict["gst_rate"] = _parse_numeric(payload.gst_rate)
        if payload.budget_percentage is not None:
            update_dict["budget_percentage"] = _parse_numeric(payload.budget_percentage)
        if payload.budget_amount is not None:
            update_dict["budget_amount"] = _parse_numeric(payload.budget_amount)
        if payload.status is not None:
            update_dict["status"] = payload.status.strip() if payload.status else "Active"

        updated = self.repo.update(item_id, update_dict)
        return _format_gbpa_response(updated)

    def delete_gbpa(self, item_id: int) -> Dict[str, Any]:
        success = self.repo.delete(item_id)
        if not success:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"GBPA Item with ID {item_id} not found."
            )
        return {"success": True, "message": f"GBPA Item {item_id} deleted successfully."}

    def process_bulk_upload(self, file_content: bytes, filename: str = "") -> Dict[str, Any]:
        """Parses Excel or CSV content and bulk-inserts records into indus_customer_gbpa."""
        import io
        import csv
        import re

        if not file_content:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded file is empty."
            )

        rows_dict_list = []
        raw_headers = []
        is_excel = filename.lower().endswith((".xlsx", ".xls")) or file_content.startswith(b"PK\x03\x04")

        if is_excel:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
                sheet = wb.active
                if sheet is None:
                    raise ValueError("Excel file contains no active sheet.")
                
                for col in range(1, sheet.max_column + 1):
                    val = sheet.cell(1, col).value
                    raw_headers.append(str(val).strip() if val is not None else f"col_{col}")

                for row_idx in range(2, sheet.max_row + 1):
                    row_data = {}
                    has_data = False
                    for col_idx, header in enumerate(raw_headers, start=1):
                        cell_val = sheet.cell(row_idx, col_idx).value
                        if cell_val is not None and str(cell_val).strip():
                            has_data = True
                        row_data[header] = cell_val
                    if has_data:
                        rows_dict_list.append(row_data)
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Invalid Excel file or corrupted workbook: {str(e)}"
                )
        else:
            try:
                try:
                    decoded = file_content.decode("utf-8-sig")
                except UnicodeDecodeError:
                    decoded = file_content.decode("latin-1")
                
                reader = csv.DictReader(io.StringIO(decoded))
                raw_headers = reader.fieldnames or []
                for r in reader:
                    if any(v and str(v).strip() for v in r.values()):
                        rows_dict_list.append(r)
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Failed to read CSV file: {str(e)}"
                )

        def norm(h: str) -> str:
            return re.sub(r'[^a-zA-Z0-9]', '', str(h)).lower()

        norm_headers = [norm(h) for h in raw_headers if h is not None]
        # Must have item code/name and rate/hsn/sac/budget
        has_gbpa_structure = any("itemcode" in nh or "itemname" in nh for nh in norm_headers) and any("rate" in nh or "hsn" in nh or "budget" in nh for nh in norm_headers)

        if not has_gbpa_structure:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid GBPA Upload Template. Please use the official GBPA Upload Template."
            )

        if not rows_dict_list:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No GBPA data rows found in uploaded file."
            )

        new_records = []
        row_errors = []
        incoming_codes = []

        for idx, row in enumerate(rows_dict_list, start=2):
            try:
                def get_val(*keys: str, default: Any = None) -> Any:
                    for k in keys:
                        if k in row and row[k] is not None:
                            return row[k]
                        k_norm = norm(k)
                        for rk, rv in row.items():
                            if rv is not None and norm(rk) == k_norm:
                                return rv
                    return default

                cust_name = str(get_val("Customer", default="Indus Tower Ltd") or "Indus Tower Ltd").strip()
                item_code = str(get_val("Item Code", "ItemCode", default="") or "").strip()
                item_name = str(get_val("Item Name", "ItemName", default="") or "").strip()
                if not item_code and not item_name:
                    item_code = f"GBPA-{idx:04d}"
                    item_name = f"GBPA Item {idx}"
                elif not item_name:
                    item_name = item_code
                elif not item_code:
                    item_code = f"GBPA-{idx:04d}"

                item_type = str(get_val("ItemType", "Item Type", default="Capex") or "Capex").strip()
                item_desc = str(get_val("Item Description", "Description", default="") or "").strip() or item_name
                uom = str(get_val("Uom", "UOM", default="Pcs") or "Pcs").strip()
                
                rate_val = _parse_numeric(get_val("Rate", default="0"))
                hsn_sac = str(get_val("HSN / SAC", "HSN/SAC", "HSN_SAC", default="HSN") or "HSN").strip()
                hsn_sac_code = str(get_val("HSN / SAC Code", "HSN/SAC Code", "HSN Code", default="") or "").strip()
                
                budget_pct = _parse_numeric(get_val("Budget (%)", "Budget %", "Budget Percentage", default=None))
                budget_amt = _parse_numeric(get_val("Budget (₹)", "Budget Amount", "Budget Rs", default=None))

                gbpa_record = IndusCustomerGbpa(
                    company_name="Nexus",
                    customer_name=cust_name,
                    item_code=item_code,
                    item_name=item_name,
                    item_description=item_desc,
                    item_type=item_type,
                    hsn_sac=hsn_sac,
                    hsn_sac_code=hsn_sac_code if hsn_sac_code else None,
                    uom=uom,
                    rate=rate_val,
                    budget_percentage=budget_pct,
                    budget_amount=budget_amt,
                    status="Active"
                )
                new_records.append(gbpa_record)
                incoming_codes.append(item_code)
            except Exception as e:
                row_errors.append(f"Row {idx}: {str(e)}")

        if not new_records and row_errors:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse records: {'; '.join(row_errors[:5])}"
            )

        # Idempotency check: if all incoming item_codes already exist in DB
        if incoming_codes:
            existing_count = self.db.query(IndusCustomerGbpa).filter(
                IndusCustomerGbpa.item_code.in_(incoming_codes)
            ).count()
            if existing_count == len(incoming_codes):
                return {
                    "success": True,
                    "imported_count": len(incoming_codes),
                    "is_retry": True,
                    "message": "Upload already processed. No duplicate records created.",
                    "errors": []
                }

        try:
            self.repo.bulk_create(new_records)
        except Exception as e:
            self.db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Database error during bulk insert: {str(e)}"
            )

        return {
            "success": True,
            "imported_count": len(new_records),
            "errors": row_errors
        }


