from typing import Optional, List, Dict, Any
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError

from app.models.indus_site import IndusSiteDetails
from app.repositories.indus_site_repository import IndusSiteRepository
from app.schemas.indus_site import (
    IndusSiteCreate,
    IndusSiteUpdate,
    IndusSiteResponse,
    IndusSiteListResponse,
    IndusSiteContactItem,
    IndusSiteContactSave,
    IndusSiteContactListResponse,
    parse_status_field
)

class IndusSiteService:
    def __init__(self, db: Session):
        self.db = db
        self.repo = IndusSiteRepository(db)

    def _extract_contacts(self, site: IndusSiteDetails) -> List[IndusSiteContactItem]:
        contacts: List[IndusSiteContactItem] = []
        if site.fse_name or site.fse_contact_number or site.fse_email:
            contacts.append(
                IndusSiteContactItem(
                    id="SC-FSE",
                    name=site.fse_name or "",
                    designation="Site Engineer",
                    contact=site.fse_contact_number or "",
                    contact_number=site.fse_contact_number or "",
                    email=site.fse_email or "",
                    status=site.status or "Active",
                    site_id=site.site_id,
                    site_code=site.site_code or ""
                )
            )
        if site.aom_name or site.aom_contact_number or site.aom_email:
            contacts.append(
                IndusSiteContactItem(
                    id="SC-AOM",
                    name=site.aom_name or "",
                    designation="Operations Officer",
                    contact=site.aom_contact_number or "",
                    contact_number=site.aom_contact_number or "",
                    email=site.aom_email or "",
                    status=site.status or "Active",
                    site_id=site.site_id,
                    site_code=site.site_code or ""
                )
            )
        return contacts

    def _to_response(self, site: IndusSiteDetails) -> IndusSiteResponse:
        contacts = self._extract_contacts(site)
        return IndusSiteResponse(
            id=site.site_id,
            site_id=site.site_id,
            site_code=site.site_code,
            siteCode=site.site_code,
            siteId=site.site_code,
            wh_id=site.wh_id,
            whId=site.wh_id,
            site_name=site.site_name,
            siteName=site.site_name,
            tower_type=site.tower_type,
            towerType=site.tower_type,
            district=site.district,
            town=site.town,
            address=site.address,
            latitude=str(site.latitude) if site.latitude is not None else None,
            longitude=str(site.longitude) if site.longitude is not None else None,
            transport_zone=site.transport_zone,
            transportZone=site.transport_zone,
            fse_name=site.fse_name,
            fse_contact_number=site.fse_contact_number,
            fse_email=site.fse_email,
            aom_name=site.aom_name,
            aom_contact_number=site.aom_contact_number,
            aom_email=site.aom_email,
            status=site.status or "Active",
            company_name=site.company_name or "Nexus",
            customer_name=site.customer_name or "Indus Tower Ltd",
            contacts=contacts,
            created_at=site.created_at,
            updated_at=site.updated_at
        )

    def get_sites_list(
        self,
        page: int = 1,
        page_size: int = 50,
        search: Optional[str] = None,
        tower_type: Optional[str] = None,
        status_filter: Optional[str] = None,
        sort_by: str = "site_id",
        sort_desc: bool = True
    ) -> IndusSiteListResponse:
        items, total = self.repo.list_sites(
            page=page,
            page_size=page_size,
            search=search,
            tower_type=tower_type,
            status_filter=status_filter,
            sort_by=sort_by,
            sort_desc=sort_desc
        )
        return IndusSiteListResponse(
            total=total,
            items=[self._to_response(item) for item in items],
            page=page,
            page_size=page_size
        )

    def get_site_by_id(self, site_id: int) -> IndusSiteResponse:
        site = self.repo.get_by_id(site_id)
        if not site:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Site with ID {site_id} not found."
            )
        return self._to_response(site)

    def create_site(self, payload: IndusSiteCreate) -> IndusSiteResponse:
        site_code = (payload.site_code or "").strip()
        site_name = (payload.site_name or "").strip()

        if not site_code:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Site ID is required."
            )
        if not site_name:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Site Name is required."
            )

        # Enforce unique site_code
        existing_code = self.repo.get_by_code(site_code)
        if existing_code:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"Site ID '{site_code}' already exists."
            )

        # Enforce unique site_name
        existing_name = self.repo.get_by_name(site_name)
        if existing_name:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"Site Name '{site_name}' already exists."
            )

        site_data: Dict[str, Any] = {
            "company_name": payload.company_name or "Nexus",
            "customer_name": payload.customer_name or "Indus Tower Ltd",
            "site_code": site_code,
            "wh_id": payload.wh_id.strip() if payload.wh_id else None,
            "site_name": site_name,
            "tower_type": payload.tower_type.strip() if payload.tower_type else None,
            "district": payload.district.strip() if payload.district else None,
            "town": payload.town.strip() if payload.town else None,
            "address": payload.address.strip() if payload.address else None,
            "latitude": payload.latitude if payload.latitude not in ("", None) else None,
            "longitude": payload.longitude if payload.longitude not in ("", None) else None,
            "transport_zone": payload.transport_zone.strip() if payload.transport_zone else None,
            "fse_name": payload.fse_name.strip() if payload.fse_name else None,
            "fse_contact_number": payload.fse_contact_number.strip() if payload.fse_contact_number else None,
            "fse_email": payload.fse_email.strip() if payload.fse_email else None,
            "aom_name": payload.aom_name.strip() if payload.aom_name else None,
            "aom_contact_number": payload.aom_contact_number.strip() if payload.aom_contact_number else None,
            "aom_email": payload.aom_email.strip() if payload.aom_email else None,
            "status": parse_status_field(payload.status)
        }

        try:
            site = self.repo.create(site_data)
            return self._to_response(site)
        except IntegrityError as e:
            self.db.rollback()
            err_msg = str(e.orig).lower()
            if "site_code" in err_msg or "site_id" in err_msg:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"Site ID '{site_code}' already exists."
                )
            if "site_name" in err_msg:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"Site Name '{site_name}' already exists."
                )
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"Database constraint violation: {str(e.orig)}"
            )

    def update_site(self, site_id: int, payload: IndusSiteUpdate) -> IndusSiteResponse:
        existing = self.repo.get_by_id(site_id)
        if not existing:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Site with ID {site_id} not found."
            )

        site_code = (payload.site_code or existing.site_code or "").strip()
        site_name = (payload.site_name or existing.site_name or "").strip()

        if not site_code:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Site ID cannot be empty."
            )
        if not site_name:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Site Name cannot be empty."
            )

        # Check unique site_code (excluding current site_id)
        existing_code = self.repo.get_by_code(site_code, exclude_id=site_id)
        if existing_code:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"Site ID '{site_code}' already exists on another site."
            )

        # Check unique site_name (excluding current site_id)
        existing_name = self.repo.get_by_name(site_name, exclude_id=site_id)
        if existing_name:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"Site Name '{site_name}' already exists on another site."
            )

        update_dict: Dict[str, Any] = {}
        if payload.site_code is not None:
            update_dict["site_code"] = site_code
        if payload.site_name is not None:
            update_dict["site_name"] = site_name
        if payload.wh_id is not None:
            update_dict["wh_id"] = payload.wh_id.strip() if payload.wh_id else None
        if payload.tower_type is not None:
            update_dict["tower_type"] = payload.tower_type.strip() if payload.tower_type else None
        if payload.district is not None:
            update_dict["district"] = payload.district.strip() if payload.district else None
        if payload.town is not None:
            update_dict["town"] = payload.town.strip() if payload.town else None
        if payload.address is not None:
            update_dict["address"] = payload.address.strip() if payload.address else None
        if payload.latitude is not None:
            update_dict["latitude"] = payload.latitude if payload.latitude not in ("", None) else None
        if payload.longitude is not None:
            update_dict["longitude"] = payload.longitude if payload.longitude not in ("", None) else None
        if payload.transport_zone is not None:
            update_dict["transport_zone"] = payload.transport_zone.strip() if payload.transport_zone else None
        if payload.fse_name is not None:
            update_dict["fse_name"] = payload.fse_name.strip() if payload.fse_name else None
        if payload.fse_contact_number is not None:
            update_dict["fse_contact_number"] = payload.fse_contact_number.strip() if payload.fse_contact_number else None
        if payload.fse_email is not None:
            update_dict["fse_email"] = payload.fse_email.strip() if payload.fse_email else None
        if payload.aom_name is not None:
            update_dict["aom_name"] = payload.aom_name.strip() if payload.aom_name else None
        if payload.aom_contact_number is not None:
            update_dict["aom_contact_number"] = payload.aom_contact_number.strip() if payload.aom_contact_number else None
        if payload.aom_email is not None:
            update_dict["aom_email"] = payload.aom_email.strip() if payload.aom_email else None
        if payload.status is not None:
            update_dict["status"] = parse_status_field(payload.status)

        try:
            updated = self.repo.update(site_id, update_dict)
            return self._to_response(updated)
        except IntegrityError as e:
            self.db.rollback()
            err_msg = str(e.orig).lower()
            if "site_code" in err_msg or "site_id" in err_msg:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"Site ID '{site_code}' already exists."
                )
            if "site_name" in err_msg:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail=f"Site Name '{site_name}' already exists."
                )
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"Database constraint violation: {str(e.orig)}"
            )

    def delete_site(self, site_id: int) -> dict:
        site = self.repo.get_by_id(site_id)
        if not site:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Site with ID {site_id} not found."
            )
        self.repo.delete(site_id)
        return {"success": True, "message": f"Site {site_id} deleted successfully."}

    def get_site_contacts(self, site_id: int) -> IndusSiteContactListResponse:
        site = self.repo.get_by_id(site_id)
        if not site:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Site with ID {site_id} not found."
            )
        contacts = self._extract_contacts(site)
        return IndusSiteContactListResponse(
            total=len(contacts),
            items=contacts
        )

    def save_site_contact(self, site_id: int, payload: IndusSiteContactSave) -> IndusSiteContactListResponse:
        site = self.repo.get_by_id(site_id)
        if not site:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Site with ID {site_id} not found."
            )

        desig = str(payload.designation or "").lower()
        cid = str(payload.id or "").upper()
        
        update_dict: Dict[str, Any] = {}
        if "engineer" in desig or "fse" in desig or cid == "SC-FSE" or cid == "1":
            if payload.name is not None:
                update_dict["fse_name"] = payload.name.strip() if payload.name else None
            if payload.contact_number is not None:
                update_dict["fse_contact_number"] = payload.contact_number.strip() if payload.contact_number else None
            if payload.email is not None:
                update_dict["fse_email"] = payload.email.strip() if payload.email else None
        else: # Default to AOM / Operations Officer
            if payload.name is not None:
                update_dict["aom_name"] = payload.name.strip() if payload.name else None
            if payload.contact_number is not None:
                update_dict["aom_contact_number"] = payload.contact_number.strip() if payload.contact_number else None
            if payload.email is not None:
                update_dict["aom_email"] = payload.email.strip() if payload.email else None

        updated = self.repo.update(site_id, update_dict)
        contacts = self._extract_contacts(updated)
        return IndusSiteContactListResponse(
            total=len(contacts),
            items=contacts
        )

    def process_bulk_upload(self, file_content: bytes, filename: str = "") -> Dict[str, Any]:
        """Parses Excel or CSV content and bulk-inserts or updates records in indus_site_details."""
        import io
        import csv
        import re

        if not file_content:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded file is empty."
            )

        rows_dict_list = []
        raw_headers = []
        is_excel = filename.lower().endswith((".xlsx", ".xls")) or file_content.startswith(b"PK\x03\x04")

        if is_excel:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
                sheet = wb.active
                if sheet is None:
                    raise ValueError("Excel file contains no active sheet.")
                
                for col in range(1, sheet.max_column + 1):
                    val = sheet.cell(1, col).value
                    raw_headers.append(str(val).strip() if val is not None else f"col_{col}")

                for row_idx in range(2, sheet.max_row + 1):
                    row_data = {}
                    has_data = False
                    for col_idx, header in enumerate(raw_headers, start=1):
                        cell_val = sheet.cell(row_idx, col_idx).value
                        if cell_val is not None and str(cell_val).strip():
                            has_data = True
                        row_data[header] = cell_val
                    if has_data:
                        rows_dict_list.append(row_data)
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Invalid Excel file or corrupted workbook: {str(e)}"
                )
        else:
            try:
                try:
                    decoded = file_content.decode("utf-8-sig")
                except UnicodeDecodeError:
                    decoded = file_content.decode("latin-1")
                
                reader = csv.DictReader(io.StringIO(decoded))
                raw_headers = reader.fieldnames or []
                for r in reader:
                    if any(v and str(v).strip() for v in r.values()):
                        rows_dict_list.append(r)
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Failed to read CSV file: {str(e)}"
                )

        def norm(h: str) -> str:
            return re.sub(r'[^a-zA-Z0-9]', '', str(h)).lower()

        norm_headers = [norm(h) for h in raw_headers if h is not None]
        # Must have site id / site code or site name or fse / aom
        has_site_structure = any("siteid" in nh or "sitecode" in nh or "sitename" in nh for nh in norm_headers) and any("towertype" in nh or "district" in nh or "town" in nh or "fse" in nh or "aom" in nh or "transport" in nh for nh in norm_headers)

        if not has_site_structure:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid Site Upload Template. Please use the official Site Upload Template."
            )

        if not rows_dict_list:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No Site data rows found in uploaded file."
            )

        new_records = []
        row_errors = []
        incoming_codes = []

        for idx, row in enumerate(rows_dict_list, start=2):
            try:
                def get_val(*keys: str, default: Any = None) -> Any:
                    for k in keys:
                        if k in row and row[k] is not None:
                            return row[k]
                        k_norm = norm(k)
                        for rk, rv in row.items():
                            if rv is not None and norm(rk) == k_norm:
                                return rv
                    return default

                cust_name = str(get_val("Customer", default="Indus Tower Ltd") or "Indus Tower Ltd").strip()
                site_code = str(get_val("Site ID", "SiteID", "Site Code", default="") or "").strip()
                wh_id = str(get_val("WH ID", "WHID", "Warehouse ID", default="") or "").strip()
                site_name = str(get_val("Site Name", "SiteName", default="") or "").strip()
                
                if not site_code and not site_name:
                    site_code = f"SITE-{idx:04d}"
                    site_name = f"Site {idx}"
                elif not site_name:
                    site_name = f"Site {site_code}"
                elif not site_code:
                    site_code = f"SITE-{idx:04d}"

                tower_type = str(get_val("Tower Type", "TowerType", default="GBT") or "GBT").strip()
                district = str(get_val("District", default="") or "").strip()
                town = str(get_val("Town", default="") or "").strip()
                address = str(get_val("Site Address", "Address", default="") or "").strip()
                
                lat_raw = get_val("Lattitude", "Latitude", default=None)
                long_raw = get_val("Longtitude", "Longitude", default=None)
                lat_val = None
                long_val = None
                if lat_raw is not None and str(lat_raw).strip():
                    try:
                        lat_val = float(str(lat_raw).strip())
                    except Exception:
                        pass
                if long_raw is not None and str(long_raw).strip():
                    try:
                        long_val = float(str(long_raw).strip())
                    except Exception:
                        pass

                transport_zone = str(get_val("Transport Zone", "TransportZone", default="") or "").strip()
                fse_name = str(get_val("FSE Name", "FSEName", "FSE", default="") or "").strip()
                aom_name = str(get_val("AOM Name", "AOMName", "AOM", default="") or "").strip()

                # Check if site already exists in database
                existing_site = self.repo.get_by_code(site_code)
                if existing_site:
                    # Update contact or fields
                    if fse_name:
                        existing_site.fse_name = fse_name
                    if aom_name:
                        existing_site.aom_name = aom_name
                    if wh_id:
                        existing_site.wh_id = wh_id
                    if tower_type:
                        existing_site.tower_type = tower_type
                    if district:
                        existing_site.district = district
                    if town:
                        existing_site.town = town
                    if address:
                        existing_site.address = address
                    if lat_val is not None:
                        existing_site.latitude = lat_val
                    if long_val is not None:
                        existing_site.longitude = long_val
                    if transport_zone:
                        existing_site.transport_zone = transport_zone
                    self.db.commit()
                else:
                    site_record = IndusSiteDetails(
                        company_name="Nexus",
                        customer_name=cust_name,
                        site_code=site_code,
                        wh_id=wh_id if wh_id else None,
                        site_name=site_name,
                        tower_type=tower_type,
                        district=district if district else None,
                        town=town if town else None,
                        address=address if address else None,
                        latitude=lat_val,
                        longitude=long_val,
                        transport_zone=transport_zone if transport_zone else None,
                        fse_name=fse_name if fse_name else None,
                        aom_name=aom_name if aom_name else None,
                        status="Active"
                    )
                    new_records.append(site_record)

                incoming_codes.append(site_code)
            except Exception as e:
                row_errors.append(f"Row {idx}: {str(e)}")

        if not new_records and not incoming_codes and row_errors:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse records: {'; '.join(row_errors[:5])}"
            )

        if new_records:
            try:
                self.repo.bulk_create(new_records)
            except Exception as e:
                self.db.rollback()
                raise HTTPException(
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail=f"Database error during bulk insert: {str(e)}"
                )

        return {
            "success": True,
            "imported_count": len(incoming_codes),
            "errors": row_errors
        }

