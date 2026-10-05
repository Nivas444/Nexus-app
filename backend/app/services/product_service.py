import io
import csv
from typing import List, Optional, Tuple, Dict, Any
from decimal import Decimal
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError

from app.models.product import CompanyProduct
from app.repositories.product_repository import ProductRepository
from app.schemas.product import (
    ProductCreate,
    ProductUpdate,
    ProductStatusUpdate,
    ProductResponse,
    ProductListResponse,
    parse_numeric_field,
    parse_status_field
)

class ProductService:
    def __init__(self, db: Session):
        self.repository = ProductRepository(db)

    def _to_response_dto(self, prod: CompanyProduct) -> ProductResponse:
        """Converts an ORM CompanyProduct model instance to a ProductResponse schema with dual field mappings."""
        def _fmt_clean(val: Any) -> str:
            if val is None:
                return ""
            try:
                f = float(val)
                if f.is_integer():
                    return str(int(f))
                return f"{f:.2f}".rstrip("0").rstrip(".")
            except Exception:
                return str(val)

        gst_num = float(prod.gst_rate) if prod.gst_rate is not None else 18.0
        gst_str = f"{int(gst_num)}%" if gst_num.is_integer() else f"{gst_num}%"

        oh_num = float(prod.oh) if prod.oh is not None else 2.0
        oh_str = f"{int(oh_num)}%" if oh_num.is_integer() else f"{oh_num}%"

        margin_num = float(prod.profit) if prod.profit is not None else 5.0
        margin_str = f"{int(margin_num)}%" if margin_num.is_integer() else f"{margin_num}%"

        msq_val = _fmt_clean(prod.msq)
        moq_val = _fmt_clean(prod.moq)

        cmp_val = float(prod.cmp) if prod.cmp is not None else 450000.00
        formatted_price = f"{cmp_val:,.2f}"

        return ProductResponse(
            id=prod.material_id,
            material_id=prod.material_id,
            product_name=prod.product_name or "",
            productName=prod.product_name or "",
            material_head=prod.material_head or prod.product_name or "",
            productHead=prod.material_head or prod.product_name or "",
            material_category=prod.material_category or "Tower Infrastructure",
            productCategory=prod.material_category or "Tower Infrastructure",
            material_code=prod.material_code or "",
            productCode=prod.material_code or "",
            hsn_code=prod.hsn_code or "",
            hsnCode=prod.hsn_code or "",
            material_description=prod.material_description or prod.product_name or "",
            make=prod.make or "",
            uom=prod.uom or "Nos",
            sale_uom=prod.sale_uom or prod.uom or "Nos",
            saleUom=prod.sale_uom or prod.uom or "Nos",
            ucf=float(prod.ucf) if prod.ucf is not None else None,
            msq=msq_val,
            moq=moq_val,
            oh=oh_str,
            margin=margin_str,
            profit=margin_num,
            gst_rate=gst_num,
            gstRate=gst_str,
            gst=gst_str,
            price=formatted_price,
            stockPrices=formatted_price,
            status=prod.status or "Active",
            company_name=prod.company_name or "Nexus",
            created_at=prod.created_at,
            updated_at=prod.updated_at
        )

    def get_products_list(
        self,
        page: int = 1,
        page_size: int = 100,
        search: Optional[str] = None,
        category: Optional[str] = None,
        head: Optional[str] = None,
        status_filter: Optional[str] = None,
        sort_by: str = "material_id",
        sort_desc: bool = True
    ) -> ProductListResponse:
        page = max(1, page)
        page_size = max(1, min(page_size, 500))
        skip = (page - 1) * page_size

        items, total = self.repository.get_list(
            skip=skip,
            limit=page_size,
            search=search,
            category=category,
            head=head,
            status=status_filter,
            sort_by=sort_by,
            sort_desc=sort_desc
        )

        response_items = [self._to_response_dto(item) for item in items]
        return ProductListResponse(
            total=total,
            items=response_items,
            page=page,
            page_size=page_size
        )

    def get_product_by_id(self, material_id: int) -> ProductResponse:
        prod = self.repository.get_by_id(material_id)
        if not prod:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Product with ID {material_id} not found"
            )
        return self._to_response_dto(prod)

    def create_product(self, data: ProductCreate) -> ProductResponse:
        p_name = (data.product_name or data.productName or "").strip()
        if not p_name:
            raise HTTPException(
                status_code=422,
                detail="Product Name is required."
            )

        # Check duplicate product name
        if self.repository.get_by_product_name(p_name):
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Product name already exists."
            )

        m_head = (data.material_head or data.productHead or p_name).strip()
        # Check duplicate material head
        if self.repository.get_by_material_head(m_head):
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Material head already exists."
            )

        m_code = (data.material_code or data.productCode or f"PRD-{abs(hash(p_name)) % 10000:04d}").strip()
        if data.material_code and self.repository.get_by_material_code(m_code):
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Material code already exists."
            )

        category = (data.material_category or data.productCategory or "Tower Infrastructure").strip()
        hsn = (data.hsn_code or data.hsnCode or "73082019").strip()
        uom_val = (data.uom or "Nos").strip()
        sale_uom_val = (data.sale_uom or data.saleUom or uom_val).strip()

        ucf_val = parse_numeric_field(data.ucf)
        msq_val = parse_numeric_field(data.msq)
        moq_val = parse_numeric_field(data.moq)
        oh_val = parse_numeric_field(data.oh) or Decimal("2.00")
        profit_val = parse_numeric_field(data.profit or data.margin) or Decimal("5.00")
        gst_val = parse_numeric_field(data.gst_rate or data.gstRate or data.gst) or Decimal("18.00")
        status_val = parse_status_field(data.status)

        new_product = CompanyProduct(
            product_name=p_name,
            material_head=m_head,
            material_category=category,
            material_code=m_code,
            hsn_code=hsn,
            material_description=data.material_description or p_name,
            make=data.make or "Standard",
            uom=uom_val,
            sale_uom=sale_uom_val,
            ucf=ucf_val,
            msq=msq_val,
            moq=moq_val,
            oh=oh_val,
            profit=profit_val,
            gst_rate=gst_val,
            status=status_val,
            company_name=data.company_name or "Nexus"
        )

        try:
            saved = self.repository.create(new_product)
            return self._to_response_dto(saved)
        except IntegrityError as e:
            self.repository.db.rollback()
            err_str = str(e.orig).lower() if hasattr(e, 'orig') else str(e).lower()
            if "product_name" in err_str:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Product name already exists."
                )
            elif "material_head" in err_str:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Material head already exists."
                )
            elif "material_code" in err_str:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Material code already exists."
                )
            else:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Product record already exists."
                )

    def update_product(self, material_id: int, data: ProductUpdate) -> ProductResponse:
        prod = self.repository.get_by_id(material_id)
        if not prod:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Product with ID {material_id} not found"
            )

        # ─── STRICT EDIT RESTRICTIONS ENFORCEMENT ───────────────────────────
        # HSN code, Product name, and Material head MUST NOT be modified in edit mode
        req_pname = data.product_name or data.productName
        if req_pname is not None and req_pname.strip() and req_pname.strip() != prod.product_name:
            raise HTTPException(
                status_code=422,
                detail="Product name cannot be modified in edit mode."
            )

        req_mhead = data.material_head or data.productHead
        if req_mhead is not None and req_mhead.strip() and req_mhead.strip() != prod.material_head:
            raise HTTPException(
                status_code=422,
                detail="Material head cannot be modified in edit mode."
            )

        req_hsn = data.hsn_code or data.hsnCode
        if req_hsn is not None and req_hsn.strip() and req_hsn.strip() != prod.hsn_code:
            raise HTTPException(
                status_code=422,
                detail="HSN code cannot be modified in edit mode."
            )

        # ─── UPDATE ONLY PERMITTED FIELDS ──────────────────────────────────
        # 1. GST Rate
        gst_input = data.gst_rate if data.gst_rate is not None else (data.gstRate if data.gstRate is not None else data.gst)
        if gst_input is not None:
            parsed_gst = parse_numeric_field(gst_input)
            if parsed_gst is not None:
                prod.gst_rate = parsed_gst

        # 2. MSQ
        if data.msq is not None:
            prod.msq = parse_numeric_field(data.msq)

        # 3. MOQ
        if data.moq is not None:
            prod.moq = parse_numeric_field(data.moq)

        # 4. Margin (profit)
        margin_input = data.profit if data.profit is not None else data.margin
        if margin_input is not None:
            parsed_margin = parse_numeric_field(margin_input)
            if parsed_margin is not None:
                prod.profit = parsed_margin

        # 5. OH
        if data.oh is not None:
            parsed_oh = parse_numeric_field(data.oh)
            if parsed_oh is not None:
                prod.oh = parsed_oh

        # 6. Status
        if data.status is not None:
            prod.status = parse_status_field(data.status)

        try:
            updated = self.repository.update(prod)
            return self._to_response_dto(updated)
        except IntegrityError as e:
            self.repository.db.rollback()
            err_str = str(e.orig).lower() if hasattr(e, 'orig') else str(e).lower()
            if "product_name" in err_str:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Product name already exists."
                )
            elif "material_head" in err_str:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Material head already exists."
                )
            else:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail="Product update failed due to constraint violation."
                )

    def update_status(self, material_id: int, status_data: ProductStatusUpdate) -> ProductResponse:
        prod = self.repository.get_by_id(material_id)
        if not prod:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Product with ID {material_id} not found"
            )

        if status_data.is_active is not None:
            prod.status = "Active" if status_data.is_active else "In - Active"
        elif status_data.status is not None:
            prod.status = parse_status_field(status_data.status)

        updated = self.repository.update(prod)
        return self._to_response_dto(updated)

    def delete_product(self, material_id: int) -> Dict[str, Any]:
        prod = self.repository.get_by_id(material_id)
        if not prod:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Product with ID {material_id} not found"
            )
        self.repository.delete(prod)
        return {"success": True, "message": f"Product {material_id} deleted successfully."}

    def process_bulk_upload(self, file_content: bytes, filename: str = "") -> Dict[str, Any]:
        """Parses Excel or CSV content and bulk-inserts records into company_products."""
        if not file_content:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded file is empty."
            )

        rows_dict_list: List[Dict[str, Any]] = []
        raw_headers: List[str] = []
        is_excel = filename.lower().endswith((".xlsx", ".xls")) or file_content.startswith(b"PK\x03\x04")

        if is_excel:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
                sheet = wb.active
                if sheet is None:
                    raise ValueError("Excel file contains no active sheet.")
                
                for col in range(1, sheet.max_column + 1):
                    val = sheet.cell(1, col).value
                    raw_headers.append(str(val).strip() if val is not None else f"col_{col}")

                for row_idx in range(2, sheet.max_row + 1):
                    row_data = {}
                    has_data = False
                    for col_idx, header in enumerate(raw_headers, start=1):
                        cell_val = sheet.cell(row_idx, col_idx).value
                        if cell_val is not None and str(cell_val).strip():
                            has_data = True
                        row_data[header] = cell_val
                    if has_data:
                        rows_dict_list.append(row_data)
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Invalid Excel file or corrupted workbook: {str(e)}"
                )
        else:
            try:
                try:
                    decoded = file_content.decode("utf-8-sig")
                except UnicodeDecodeError:
                    decoded = file_content.decode("latin-1")
                
                reader = csv.DictReader(io.StringIO(decoded))
                raw_headers = reader.fieldnames or []
                for r in reader:
                    if any(v and str(v).strip() for v in r.values()):
                        rows_dict_list.append(r)
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Failed to read CSV file: {str(e)}"
                )

        # Template Header Validation
        import re
        def norm(h: str) -> str:
            return re.sub(r'[^a-zA-Z0-9]', '', str(h)).lower()

        norm_headers = [norm(h) for h in raw_headers if h is not None]
        has_material_or_product = any("material" in nh or "product" in nh for nh in norm_headers)

        if not has_material_or_product:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid Products Upload Template. Please use the official Products Upload Template."
            )

        if not rows_dict_list:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No product data rows found in uploaded file."
            )

        new_records: List[CompanyProduct] = []
        row_errors: List[str] = []
        incoming_heads: List[str] = []

        for idx, row in enumerate(rows_dict_list, start=2):
            try:
                def get_val(*keys: str, default: Any = None) -> Any:
                    for k in keys:
                        if k in row and row[k] is not None:
                            return row[k]
                        k_norm = norm(k)
                        for rk, rv in row.items():
                            if rv is not None and norm(rk) == k_norm:
                                return rv
                    return default

                mat_head = str(get_val("Material Head", "Product Head", "Head", "Name", default="") or "").strip()
                if not mat_head:
                    mat_head = f"Product Item {idx}"

                category = str(get_val("Material Categry", "Material Category", "Category", default="Tower Infrastructure") or "Tower Infrastructure").strip()
                mat_code = str(get_val("Material Code", "Product Code", "Code", default="") or "").strip() or f"PRD-{idx:04d}"
                mat_desc = str(get_val("Material Description", "Description", default="") or "").strip() or mat_head
                uom = str(get_val("Uom", "UOM", default="Nos") or "Nos").strip()
                sale_uom = str(get_val("Sale Uom", "Sale UOM", default="Nos") or "Nos").strip()
                
                gst_val = parse_numeric_field(get_val("GST Rate", "GST", default="18%")) or Decimal("18.00")
                ucf_val = parse_numeric_field(get_val("UCF", default="1.00")) or Decimal("1.00")
                msq_val = parse_numeric_field(get_val("MSQ", default="100")) or Decimal("100")
                moq_val = parse_numeric_field(get_val("MOQ", default="50")) or Decimal("50")
                margin_val = parse_numeric_field(get_val("Margin", default="5%")) or Decimal("5.00")
                oh_val = parse_numeric_field(get_val("OH", default="2%")) or Decimal("2.00")
                status_val = parse_status_field(get_val("Status", default="Active"))

                prod_record = CompanyProduct(
                    product_name=mat_head,
                    material_head=mat_head,
                    material_category=category,
                    material_code=mat_code,
                    material_description=mat_desc,
                    uom=uom,
                    sale_uom=sale_uom,
                    gst_rate=gst_val,
                    ucf=ucf_val,
                    msq=msq_val,
                    moq=moq_val,
                    profit=margin_val,
                    oh=oh_val,
                    status=status_val,
                    company_name="Nexus"
                )
                new_records.append(prod_record)
                incoming_heads.append(mat_head)
            except Exception as e:
                row_errors.append(f"Row {idx}: {str(e)}")

        if not new_records and row_errors:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse records: {'; '.join(row_errors[:5])}"
            )

        # Idempotency check: if all incoming material heads already exist in DB
        if incoming_heads:
            existing_count = self.repository.db.query(CompanyProduct).filter(
                CompanyProduct.material_head.in_(incoming_heads)
            ).count()
            if existing_count == len(incoming_heads):
                return {
                    "success": True,
                    "imported_count": len(incoming_heads),
                    "is_retry": True,
                    "message": "Upload already processed. No duplicate records created.",
                    "errors": []
                }

        try:
            self.repository.bulk_create(new_records)
        except Exception as e:
            self.repository.db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Database error during bulk insert: {str(e)}"
            )

        return {
            "success": True,
            "imported_count": len(new_records),
            "errors": row_errors
        }
